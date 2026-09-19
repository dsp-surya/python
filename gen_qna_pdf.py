import os
import json
import re
import time
import glob
import logging

# Suppress root logger and all Google / HTTP client sub-loggers
logging.basicConfig(level=logging.WARNING)
logging.getLogger("google_genai").setLevel(logging.ERROR)
logging.getLogger("google").setLevel(logging.ERROR)
logging.getLogger("httpx").setLevel(logging.ERROR)

from fpdf import FPDF
from pypdf import PdfReader, PdfWriter
from google import genai
from google.genai import types
from google.genai.errors import ClientError, ServerError

# Optional dependencies for standalone DOCX scaling (Option 3)
try:
    from pdf2docx import Converter
    from docx import Document
    from docx.shared import Pt
    from docx2pdf import convert
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

# Initialize Gemini Client (reads $env:GEMINI_API_KEY)
client = genai.Client()

# List of models to rotate through sequentially on Rate Limit (429), Model Not Found (404), or High Demand (503)
MODELS = [
    "models/gemini-2.5-flash",
    "models/gemini-3.5-flash",
    "models/gemini-3.6-flash",
    "models/gemini-3.7-flash",
    "models/gemini-3.8-flash",
    "models/gemini-2.5-flash-lite",
    "models/gemini-3.5-flash-lite",
    "models/gemini-3.1-flash-lite"
]
current_model_idx = 0


def sanitize_filename(name: str) -> str:
    """Cleans up filenames while retaining Telugu script."""
    clean_name = re.sub(r'[\\/*?:"<>|]', "", name).strip()
    clean_name = re.sub(r'\s+', " ", clean_name)
    return clean_name if clean_name else "Section_Output"


def load_sections_config(config_path: str = "sections.json") -> list[dict] | None:
    """Checks for a manual 'sections.json' file first to avoid AI guessing errors."""
    if os.path.exists(config_path):
        print(f"🎯 Found manual section map '{config_path}'. Bypassing AI visual scanning!")
        with open(config_path, "r", encoding="utf-8") as f:
            return json.load(f)
    return None


def save_sections_metadata(sections_info: list[dict], temp_split_dir: str):
    """Saves fallback section headings to JSON for script restarts."""
    meta_path = os.path.join(temp_split_dir, "sections_metadata.json")
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(sections_info, f, ensure_ascii=False, indent=2)


def split_scanned_pdf(input_pdf_path: str, sections_info: list[dict], temp_output_dir: str) -> list[tuple[str, str]]:
    """Splits main PDF into section files based on page maps."""
    os.makedirs(temp_output_dir, exist_ok=True)
    reader = PdfReader(input_pdf_path)
    total_pages = len(reader.pages)
    split_data = []

    for i in range(len(sections_info)):
        start_idx = sections_info[i]["page"] - 1
        heading = sections_info[i]["heading"]
        end_idx = sections_info[i + 1]["page"] - 1 if (i + 1 < len(sections_info)) else total_pages

        if start_idx >= total_pages:
            continue

        writer = PdfWriter()
        for p_idx in range(start_idx, end_idx):
            writer.add_page(reader.pages[p_idx])

        clean_heading = sanitize_filename(heading)
        section_path = os.path.join(temp_output_dir, f"temp_section_{i + 1}.pdf")

        with open(section_path, "wb") as f:
            writer.write(f)

        split_data.append((section_path, clean_heading))

    return split_data


def call_gemini_with_retry(contents, config=None, max_retries=20, initial_delay=15):
    """
    Calls Gemini API with exponential backoff and automatic model rotation.
    Handles:
      - 404 (Not Found / Deprecated): Switch model instantly without delay.
      - 429 (Rate Limit): Switch model and wait exponential delay.
      - 503 (Server High Demand / Unavailable): Switch model and wait brief delay.
    
    Option 2: Pauses for 10s if a full cycle through all models completes and wraps around to 0.
    """
    global current_model_idx
    delay = initial_delay

    for attempt in range(1, max_retries + 1):
        active_model = MODELS[current_model_idx]
        try:
            response = client.models.generate_content(
                model=active_model,
                contents=contents,
                config=config
            )
            return response

        except ClientError as e:
            if e.code == 404:
                # Model unavailable or deprecated; immediately switch without delay
                current_model_idx = (current_model_idx + 1) % len(MODELS)
                next_model = MODELS[current_model_idx]
                print(f"❌ Model {active_model} returned 404 NOT_FOUND. Switching immediately to {next_model}...")

                # Option 2: Check if cycle wrapped back to the start
                if current_model_idx == 0:
                    print("⏳ Completed full model rotation. Pausing 10s before restarting queue...")
                    time.sleep(10)

            elif e.code == 429:
                # Rate limited; switch model and apply backoff delay
                current_model_idx = (current_model_idx + 1) % len(MODELS)
                next_model = MODELS[current_model_idx]
                print(f"⚠️ Rate limit hit (429) on {active_model}.")
                print(f"   🔄 Switching to {next_model}. Waiting {delay}s before retry ({attempt}/{max_retries})...")
                time.sleep(delay)
                delay *= 2

                # Option 2: Check if cycle wrapped back to the start
                if current_model_idx == 0:
                    print("⏳ Completed full model rotation. Pausing 10s for quota recovery...")
                    time.sleep(10)

            else:
                raise e

        except ServerError as e:
            if e.code == 503:
                # High traffic / server load on this specific model; switch and try another
                current_model_idx = (current_model_idx + 1) % len(MODELS)
                next_model = MODELS[current_model_idx]
                print(f"🔥 Model {active_model} returned 503 UNAVAILABLE (High Demand).")
                print(f"   🔄 Switching to {next_model}. Waiting 5s before retry ({attempt}/{max_retries})...")
                time.sleep(5)

                # Option 2: Check if cycle wrapped back to the start
                if current_model_idx == 0:
                    print("⏳ Completed full model rotation. Pausing 10s for capacity reset...")
                    time.sleep(10)

            else:
                raise e

    raise Exception("❌ Exhausted retries across all configured models. Quotas finished, server down, or all models overloaded.")


def find_sections_with_gemini(pdf_path: str) -> list[dict]:
    """Scans visual layout to split section boundaries safely."""
    print(f"Scanning visual layout with Gemini: {os.path.basename(pdf_path)}...")

    file_ref = client.files.upload(
        file=pdf_path,
        config=types.UploadFileConfig(mime_type="application/pdf")
    )

    prompt = (
        "This PDF contains scanned handwritten Telugu notes. "
        "Examine EVERY page to find ONLY the main SECTION / LESSON START HEADINGS.\n\n"
        "CRITICAL VISUAL RULES TO IDENTIFY MAIN HEADINGS:\n"
        "1. TOP-CENTER POSITION: The heading MUST be located at the very TOP-CENTER of the page.\n"
        "2. UNDERLINE: Main headings are clearly UNDERLINED across the centered text.\n"
        "3. NO COLON: Main headings DO NOT end with a colon (:).\n"
        "4. FIRST ITEM: Topmost title on that page, placed above diagrams or bullet points (*).\n\n"
        "Task:\n"
        "Extract the EXACT Telugu text written in the top-centered main heading.\n"
        "Return strictly a JSON array of objects:\n"
        '[\n  {"page": 8, "heading": "భారత రాజ్యాంగం (Indian Polity)"}\n]'
    )

    config = types.GenerateContentConfig(response_mime_type="application/json")

    try:
        response = call_gemini_with_retry(contents=[file_ref, prompt], config=config)
    finally:
        try:
            client.files.delete(name=file_ref.name)
        except Exception:
            pass

    time.sleep(6)

    try:
        sections = json.loads(response.text)
        if not any(s.get("page") == 1 for s in sections):
            sections.insert(0, {"page": 1, "heading": "Section_1"})
        return sorted(sections, key=lambda x: x["page"])
    except Exception as e:
        print(f"⚠️ Failed to parse JSON sections: {e}")
        return [{"page": 1, "heading": "Section_1"}]


def extract_qa_page_by_page(pdf_path: str, temp_page_dir: str) -> str:
    """Extracts high-detail, essay-style Telugu Q&A page by page from any PDF file."""
    os.makedirs(temp_page_dir, exist_ok=True)
    reader = PdfReader(pdf_path)
    combined_telugu_qa = ""

    for page_idx, page in enumerate(reader.pages):
        page_num = page_idx + 1
        print(f"   -> Processing page {page_num} of {len(reader.pages)}...")

        single_page_writer = PdfWriter()
        single_page_writer.add_page(page)

        single_page_path = os.path.join(temp_page_dir, f"temp_single_page_{page_num}.pdf")
        with open(single_page_path, "wb") as f:
            single_page_writer.write(f)

        file_ref = None
        try:
            file_ref = client.files.upload(
                file=single_page_path,
                config=types.UploadFileConfig(mime_type="application/pdf")
            )

            prompt = (
                "Analyze the attached single-page scanned handwritten Telugu note thoroughly. "
                "Provide questions and answers in Telugu.\n"
                "Guidelines:\n"
                "1. Write everything strictly in clear Telugu script.\n"
                "2. Format each Q&A strictly as:\n\n"
                "ప్రశ్న: <Detailed Question in Telugu>\n"
                "సమాధానం: <Detailed Answer in Telugu>\n"
            )

            response = call_gemini_with_retry(contents=[file_ref, prompt])

            extracted_text = ""
            if response and response.text:
                extracted_text = response.text
            elif response and response.candidates and response.candidates[0].content.parts:
                extracted_text = "".join([part.text for part in response.candidates[0].content.parts if part.text])
            else:
                print(f"⚠️ Warning: Page {page_num} returned empty text. Skipping text concat.")
                extracted_text = "[ఈ పేజీ నుండి ప్రశ్నలు-సమాధానాలు సేకరించడం నిరోధించబడింది లేదా ఖాళీగా ఉంది]"

            combined_telugu_qa += f"--- పేజీ {page_num} ---\n\n" + extracted_text + "\n\n"

        finally:
            if os.path.exists(single_page_path):
                os.remove(single_page_path)
            if file_ref:
                try:
                    client.files.delete(name=file_ref.name)
                except Exception:
                    pass

        time.sleep(6.5)

    return combined_telugu_qa


def create_telugu_pdf(text_content: str, output_pdf_path: str, font_path: str = "NotoSansTelugu-Regular.ttf", scale_factor: float = 1.80):
    """
    Generates a high-readability Telugu PDF formatted specifically for standard A4 paper printing.
    Applies the scale_factor directly (Default 1.80 = 180% text scale up) to produce large, legible text natively.
    """
    if not os.path.exists(font_path):
        raise FileNotFoundError(f"Missing font file: {font_path}")

    # Calculate scaled dimensions
    base_font_size = 11.0 * scale_factor       # e.g., 11 * 1.8 = 19.8 pt
    header_font_size = 12.0 * scale_factor     # e.g., 12 * 1.8 = 21.6 pt
    line_height = 7.0 * scale_factor           # e.g., 7 * 1.8 = 12.6 mm
    header_line_height = 8.0 * scale_factor    # e.g., 8 * 1.8 = 14.4 mm

    # Standard A4 layout (210mm x 297mm)
    pdf = FPDF(orientation='P', unit='mm', format='A4')

    # Print Margins
    pdf.set_margins(left=15, top=15, right=15)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    # Enable HarfBuzz complex script shaping for proper Telugu glyph rendering
    try:
        pdf.set_text_shaping(True)
    except Exception:
        pass

    pdf.add_font("TeluguFont", style="", fname=font_path)

    # Base Font Size & Pure Black Ink
    pdf.set_font("TeluguFont", size=base_font_size)
    pdf.set_text_color(0, 0, 0)

    printable_width = pdf.epw

    lines = text_content.splitlines()

    for line in lines:
        clean_line = line.strip()
        if not clean_line:
            pdf.ln(int(4 * scale_factor))  # Scaled paragraph gap
            continue

        pdf.set_x(pdf.l_margin)

        # Format Question Line
        if clean_line.startswith("ప్రశ్న:"):
            pdf.set_font("TeluguFont", size=base_font_size)
            pdf.ln(2)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(printable_width, line_height, text=clean_line)
            pdf.ln(1)

        # Format Answer Line
        elif clean_line.startswith("సమాధానం:"):
            pdf.set_font("TeluguFont", size=base_font_size)
            pdf.multi_cell(printable_width, line_height, text=clean_line)
            pdf.ln(2)

        # Format Page/Section Markers
        elif clean_line.startswith("---"):
            pdf.set_font("TeluguFont", size=header_font_size)
            pdf.multi_cell(printable_width, header_line_height, text=clean_line)
            pdf.set_font("TeluguFont", size=base_font_size)
            pdf.ln(2)

        else:
            pdf.multi_cell(printable_width, line_height, text=clean_line)

    pdf.output(output_pdf_path)
    print(f"✅ Generated {int(scale_factor*100)}%-Scaled Telugu PDF: {output_pdf_path}\n")


def scale_docx_text(docx_path, scale_factor=1.80):
    """Opens a Word document, scales all font sizes by scale_factor, and saves back."""
    if not DOCX_SUPPORT:
        raise ImportError("pdf2docx, python-docx, and docx2pdf must be installed to use this feature.")

    doc = Document(docx_path)

    def adjust_font(run):
        if run.font.size:
            run.font.size = Pt(run.font.size.pt * scale_factor)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            adjust_font(run)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        adjust_font(run)

    doc.save(docx_path)


def process_pdfs_via_docx(input_dir, output_dir, scale_factor=1.80):
    """Converts PDFs to Word, scales text sizes, and exports back to PDF."""
    if not DOCX_SUPPORT:
        print("❌ Error: 'pdf2docx', 'python-docx', or 'docx2pdf' missing. Please install them to use Option 3.")
        return

    os.makedirs(input_dir, exist_ok=True)
    os.makedirs(output_dir, exist_ok=True)

    pdf_files = glob.glob(os.path.join(input_dir, "*.pdf"))

    if not pdf_files:
        print(f"No PDF files found in '{input_dir}' directory.")
        return

    print(f"Found {len(pdf_files)} PDF(s) to process...\n")

    for pdf_path in pdf_files:
        filename = os.path.basename(pdf_path)
        base_name = os.path.splitext(filename)[0]

        temp_docx = os.path.join(output_dir, f"{base_name}_temp.docx")
        final_pdf = os.path.join(output_dir, f"{base_name}_180pct.pdf")

        print(f"--> Processing: {filename}")

        print("    [1/3] Converting PDF to Word...")
        cv = Converter(pdf_path)
        cv.convert(temp_docx)
        cv.close()

        print(f"    [2/3] Scaling font sizes to {int(scale_factor * 100)}%...")
        scale_docx_text(temp_docx, scale_factor=scale_factor)

        print("    [3/3] Exporting back to PDF...")
        convert(temp_docx, final_pdf)

        if os.path.exists(temp_docx):
            os.remove(temp_docx)

        print(f"Done! Created '{final_pdf}'\n")

    print("All PDFs processed successfully!")


def process_single_large_scanned_pdf(input_pdf: str, temp_split_dir: str, final_qa_dir: str, scale_factor: float = 1.80):
    """Option 1: Split a single PDF into section-based Q&A PDFs."""
    os.makedirs(final_qa_dir, exist_ok=True)
    temp_page_dir = os.path.join(temp_split_dir, "single_pages")

    sections_info = load_sections_config("sections.json")

    if not sections_info:
        print("🔍 No manual 'sections.json' found. Running Gemini visual layout scan...")
        sections_info = find_sections_with_gemini(input_pdf)
        save_sections_metadata(sections_info, temp_split_dir)

    print(f"\n📌 Active Section Map ({len(sections_info)} Sections):")
    for sec in sections_info:
        print(f"   • Page {sec['page']}: {sec['heading']}")
    print("-" * 50)

    split_data = split_scanned_pdf(input_pdf, sections_info, temp_split_dir)

    for section_path, heading_name in split_data:
        output_pdf_name = f"{heading_name}.pdf"
        output_path = os.path.join(final_qa_dir, output_pdf_name)

        if os.path.exists(output_path) and os.path.getsize(output_path) > 0:
            print(f"⏩ Skipping section '{heading_name}' (Already generated)")
            continue

        print(f"\nProcessing section: '{heading_name}'...")
        combined_qa = extract_qa_page_by_page(section_path, temp_page_dir)
        create_telugu_pdf(combined_qa, output_path, scale_factor=scale_factor)


def process_directory_of_pdfs(input_dir: str, output_dir: str, temp_dir: str = "temp_batch_pages", scale_factor: float = 1.80):
    """Option 2: Process all PDFs in input_dir page-by-page and create matching output PDFs."""
    if not os.path.exists(input_dir):
        print(f"❌ Input directory '{input_dir}' does not exist.")
        return

    os.makedirs(output_dir, exist_ok=True)
    pdf_files = [f for f in os.listdir(input_dir) if f.lower().endswith(".pdf")]

    if not pdf_files:
        print(f"⚠️ No PDF files found in directory '{input_dir}'.")
        return

    print(f"📁 Found {len(pdf_files)} PDF(s) in '{input_dir}' for batch extraction.")
    print("-" * 50)

    for idx, pdf_file in enumerate(pdf_files, start=1):
        input_file_path = os.path.join(input_dir, pdf_file)
        output_file_path = os.path.join(output_dir, pdf_file)

        print(f"\n Processing file [{idx}/{len(pdf_files)}]: {pdf_file}")

        if os.path.exists(output_file_path) and os.path.getsize(output_file_path) > 0:
            print(f"⏩ Skipping '{pdf_file}' (Already generated in output directory)")
            continue

        file_temp_page_dir = os.path.join(temp_dir, f"page_cache_{idx}")
        combined_qa = extract_qa_page_by_page(input_file_path, file_temp_page_dir)
        create_telugu_pdf(combined_qa, output_file_path, scale_factor=scale_factor)


if __name__ == "__main__":
    # --- CONFIGURATION ---
    SCALE_FACTOR = 1.80  # 180% Font Scaling

    # Mode 1 Defaults (Single PDF Section Processing)
    SINGLE_INPUT_PDF = "polity_1.pdf"
    SPLIT_SECTIONS_DIR = "temp_split_sections"
    SINGLE_OUTPUT_DIR = "test"

    # Mode 2 & 3 Defaults (Directory Batch Processing)
    INPUT_BATCH_DIR = r"D:\Repos\Python\human"
    OUTPUT_BATCH_DIR = r"D:\Repos\Python\human_scale"

    print("Select Mode:")
    print("1. Section-based Q&A generation from a single PDF (Direct 180% Large-Text PDF)")
    print("2. Batch Q&A generation for all PDFs in a directory (Direct 180% Large-Text PDF)")
    print("3. External PDF-to-Word font scaling (pdf2docx -> DOCX 180% -> docx2pdf)")

    choice = input("Enter option (1, 2, or 3): ").strip()

    if choice == "1":
        process_single_large_scanned_pdf(SINGLE_INPUT_PDF, SPLIT_SECTIONS_DIR, SINGLE_OUTPUT_DIR, scale_factor=SCALE_FACTOR)
    elif choice == "2":
        process_directory_of_pdfs(INPUT_BATCH_DIR, OUTPUT_BATCH_DIR, scale_factor=SCALE_FACTOR)
    elif choice == "3":
        process_pdfs_via_docx(INPUT_BATCH_DIR, OUTPUT_BATCH_DIR, scale_factor=SCALE_FACTOR)
    else:
        print("Invalid choice selected. Exiting.")